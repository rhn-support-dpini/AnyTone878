#!/usr/bin/env python3
"""Export CPS-ready CSV files from Master_Codeplug Excel."""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from openpyxl import load_workbook

from codeplug_utils import (
    CHANNEL_TYPE_ANALOG,
    CHANNEL_TYPE_DIGITAL,
    VALID_CHANNEL_TYPES,
    default_blank_channel_row,
    find_input_file,
    find_latest_master_excel,
    is_blank_excel_row,
    is_frequency_column,
    is_single_frequency_column,
    is_numeric_index,
    join_pipe_values,
    log_message,
    normalize_frequencies_in_row,
    read_csv_dicts,
    split_pipe_values,
    today_str,
    write_csv_dicts,
)

MEMBERSHIP_COLUMNS = ("Zone Membership", "ScanList Membership")

ZONE_MEMBER_FIELDS = (
    "Zone Channel Member",
    "A Channel",
    "B Channel",
)

SCAN_MEMBER_FIELDS = (
    "Scan Channel Member",
    "Priority Channel 1",
    "Priority Channel 2",
)


def resolve_excel_path(base_dir: Path, arg_path: Optional[str]) -> Path:
    if arg_path:
        path = Path(arg_path)
        if not path.is_absolute():
            path = base_dir / path
        if not path.is_file():
            raise FileNotFoundError(f"File Excel non trovato: {path}")
        return path

    dated_matches = sorted(base_dir.glob("Master_Codeplug_*.xlsx"), reverse=True)
    if dated_matches:
        return dated_matches[0]

    latest = find_latest_master_excel(base_dir)
    if latest:
        return latest

    raise FileNotFoundError(
        "Nessun file Master_Codeplug_<AAAA-MM-GG>.xlsx trovato nella directory di lavoro"
    )


def load_metadata_files(base_dir: Path, excel_path: Path) -> Tuple[List[str], List[Dict[str, str]], List[str], List[Dict[str, str]]]:
    stem = excel_path.stem
    zone_meta = base_dir / f"{stem}_Zone.csv"
    scan_meta = base_dir / f"{stem}_ScanList.csv"

    if zone_meta.is_file() and scan_meta.is_file():
        return (*read_csv_dicts(zone_meta), *read_csv_dicts(scan_meta))

    zone_path = find_input_file(base_dir, "Zone.csv")
    scan_path = find_input_file(base_dir, "ScanList.csv")
    log_message(
        "Metadati affiancati all'Excel non trovati; uso Zone.csv e ScanList.csv originali",
        "WARNING",
    )
    return (*read_csv_dicts(zone_path), *read_csv_dicts(scan_path))


def read_excel_channels(excel_path: Path) -> Tuple[List[str], List[Dict[str, str]]]:
    workbook = load_workbook(excel_path, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]

    headers: List[str] = []
    for cell in worksheet[1]:
        headers.append(str(cell.value).strip() if cell.value is not None else "")

    rows: List[Dict[str, str]] = []
    for row_idx in range(2, worksheet.max_row + 1):
        row_data: Dict[str, str] = {}
        empty = True
        for col_idx, header in enumerate(headers, start=1):
            if not header:
                continue
            value = worksheet.cell(row=row_idx, column=col_idx).value
            if value is None:
                text = ""
            elif isinstance(value, float) and header and is_frequency_column(header):
                text = f"{value:.5f}"
            else:
                text = str(value).strip()
            if text:
                empty = False
            row_data[header] = text
        if empty and row_idx > worksheet.max_row:
            continue
        rows.append(row_data)

    workbook.close()
    return headers, rows


def classify_excel_row(
    row: Dict[str, str],
    channel_columns: List[str],
    row_number: int,
) -> str:
    """Return 'valid', 'blank', or 'anomalous'."""
    if is_blank_excel_row(row, channel_columns):
        return "blank"

    channel_type = row.get("Channel Type", "").strip()
    channel_name = row.get("Channel Name", "").strip()

    if not channel_name:
        log_message(
            f"Riga anomala in Excel (riga {row_number}): Channel Name mancante",
            "WARNING",
        )
        return "anomalous"

    if channel_type not in VALID_CHANNEL_TYPES:
        log_message(
            f"Riga anomala in Excel (riga {row_number}): Channel Type "
            f"'{channel_type or '<vuoto>'}' non riconosciuto",
            "WARNING",
        )
        return "anomalous"

    return "valid"


def process_excel_rows(
    headers: List[str],
    excel_rows: List[Dict[str, str]],
) -> Tuple[List[Dict[str, str]], Dict[str, int], List[str]]:
    channel_columns = [
        h for h in headers if h not in MEMBERSHIP_COLUMNS
    ]
    if "No." not in channel_columns and "No." in headers:
        channel_columns = headers[:]
        channel_columns = [h for h in channel_columns if h not in MEMBERSHIP_COLUMNS]

    freq_columns = [c for c in channel_columns if is_single_frequency_column(c)]
    output_rows: List[Dict[str, str]] = []
    name_to_new_index: Dict[str, int] = {}
    new_index = 0
    template_row: Optional[Dict[str, str]] = None

    for position, raw_row in enumerate(excel_rows, start=2):
        classification = classify_excel_row(raw_row, channel_columns, position)

        if classification == "blank":
            if template_row is None:
                template_row = {col: "" for col in channel_columns}
            blank_row = default_blank_channel_row(template_row)
            blank_row["No."] = ""
            output_rows.append(blank_row)
            continue

        row = {col: raw_row.get(col, "") for col in channel_columns}
        row = normalize_frequencies_in_row(
            row,
            freq_columns,
            f"Excel riga {position} '{row.get('Channel Name', '')}'",
        )

        if classification == "anomalous":
            row["No."] = ""
            output_rows.append(row)
            continue

        new_index += 1
        row["No."] = str(new_index)
        name = row.get("Channel Name", "").strip()
        if name in name_to_new_index:
            log_message(
                f"Nome canale duplicato in Excel '{name}' alla riga {position}; "
                "l'ultima occorrenza sovrascrive la mappa indici",
                "WARNING",
            )
        name_to_new_index[name] = new_index
        output_rows.append(row)
        template_row = row

    return output_rows, name_to_new_index, channel_columns


def name_to_index_value(
    name: str,
    name_to_new_index: Dict[str, int],
    context: str,
) -> str:
    name = name.strip()
    if not name:
        return name
    if is_numeric_index(name):
        return name
    if name in name_to_new_index:
        return str(name_to_new_index[name])
    log_message(
        f"Riferimento mancante: canale '{name}' non presente nel foglio Excel ({context})",
        "WARNING",
    )
    return name


def rebuild_zone_membership_from_excel(
    excel_rows: List[Dict[str, str]],
) -> Dict[str, List[str]]:
    zone_order: Dict[str, List[str]] = defaultdict(list)
    for row in excel_rows:
        name = row.get("Channel Name", "").strip()
        if not name:
            continue
        for zone_name in split_pipe_values(row.get("Zone Membership", "")):
            if name not in zone_order[zone_name]:
                zone_order[zone_name].append(name)
    return zone_order


def rebuild_scan_membership_from_excel(
    excel_rows: List[Dict[str, str]],
) -> Dict[str, List[str]]:
    scan_order: Dict[str, List[str]] = defaultdict(list)
    for row in excel_rows:
        name = row.get("Channel Name", "").strip()
        if not name:
            continue
        for scan_name in split_pipe_values(row.get("ScanList Membership", "")):
            if name not in scan_order[scan_name]:
                scan_order[scan_name].append(name)
    return scan_order


def lookup_channel_frequencies(
    channel_name: str,
    channel_rows: List[Dict[str, str]],
) -> Tuple[str, str]:
    for row in channel_rows:
        if row.get("Channel Name", "").strip() == channel_name:
            return (
                row.get("Receive Frequency", ""),
                row.get("Transmit Frequency", ""),
            )
    return "", ""


def export_zone_csv(
    zone_fieldnames: List[str],
    zone_rows: List[Dict[str, str]],
    zone_order: Dict[str, List[str]],
    name_to_new_index: Dict[str, int],
    channel_rows: List[Dict[str, str]],
) -> List[Dict[str, str]]:
    exported: List[Dict[str, str]] = []

    for zone in zone_rows:
        zone_name = zone.get("Zone Name", "").strip()
        updated = dict(zone)

        if zone_name in zone_order:
            members = zone_order[zone_name]
        else:
            members = split_pipe_values(zone.get("Zone Channel Member", ""))

        member_indices: List[str] = []
        rx_values: List[str] = []
        tx_values: List[str] = []

        for member in members:
            idx = name_to_index_value(
                member, name_to_new_index, f"Zone '{zone_name}' member"
            )
            member_indices.append(idx)
            rx, tx = lookup_channel_frequencies(member, channel_rows)
            rx_values.append(normalize_frequencies_in_row(
                {"v": rx}, ["v"], f"Zone '{zone_name}' RX"
            )["v"] if rx else "")
            tx_values.append(normalize_frequencies_in_row(
                {"v": tx}, ["v"], f"Zone '{zone_name}' TX"
            )["v"] if tx else "")

        updated["Zone Channel Member"] = join_pipe_values(member_indices)
        if "Zone Channel Member RX Frequency" in updated:
            updated["Zone Channel Member RX Frequency"] = join_pipe_values(rx_values)
        if "Zone Channel Member TX Frequency" in updated:
            updated["Zone Channel Member TX Frequency"] = join_pipe_values(tx_values)

        for field in ("A Channel", "B Channel"):
            if field in updated and updated[field].strip():
                updated[field] = name_to_index_value(
                    updated[field], name_to_new_index, f"Zone '{zone_name}'/{field}"
                )

        freq_cols = [c for c in updated if is_single_frequency_column(c)]
        updated = normalize_frequencies_in_row(updated, freq_cols, f"Zone '{zone_name}'")
        exported.append(updated)

    return exported


def export_scan_csv(
    scan_fieldnames: List[str],
    scan_rows: List[Dict[str, str]],
    scan_order: Dict[str, List[str]],
    name_to_new_index: Dict[str, int],
    channel_rows: List[Dict[str, str]],
) -> List[Dict[str, str]]:
    exported: List[Dict[str, str]] = []

    for scan in scan_rows:
        scan_name = scan.get("Scan List Name", "").strip()
        updated = dict(scan)

        if scan_name in scan_order:
            members = scan_order[scan_name]
        else:
            members = split_pipe_values(scan.get("Scan Channel Member", ""))

        member_indices: List[str] = []
        rx_values: List[str] = []
        tx_values: List[str] = []

        for member in members:
            idx = name_to_index_value(
                member, name_to_new_index, f"ScanList '{scan_name}' member"
            )
            member_indices.append(idx)
            rx, tx = lookup_channel_frequencies(member, channel_rows)
            rx_values.append(rx)
            tx_values.append(tx)

        updated["Scan Channel Member"] = join_pipe_values(member_indices)
        if "Scan Channel Member RX Frequency" in updated:
            updated["Scan Channel Member RX Frequency"] = join_pipe_values(
                [
                    normalize_frequencies_in_row({"v": v}, ["v"], f"Scan '{scan_name}'")["v"]
                    if v else ""
                    for v in rx_values
                ]
            )
        if "Scan Channel Member TX Frequency" in updated:
            updated["Scan Channel Member TX Frequency"] = join_pipe_values(
                [
                    normalize_frequencies_in_row({"v": v}, ["v"], f"Scan '{scan_name}'")["v"]
                    if v else ""
                    for v in tx_values
                ]
            )

        for field in ("Priority Channel 1", "Priority Channel 2"):
            if field in updated and updated[field].strip() and updated[field].strip() != "Off":
                updated[field] = name_to_index_value(
                    updated[field], name_to_new_index, f"ScanList '{scan_name}'/{field}"
                )

        freq_cols = [c for c in updated if is_single_frequency_column(c)]
        updated = normalize_frequencies_in_row(updated, freq_cols, f"ScanList '{scan_name}'")
        exported.append(updated)

    return exported


def main() -> int:
    base_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")
    excel_arg = sys.argv[2] if len(sys.argv) > 2 else None

    log_message("Avvio Script2_ExportCSV.py")

    try:
        excel_path = resolve_excel_path(base_dir, excel_arg)
    except FileNotFoundError as exc:
        log_message(str(exc), "ERROR")
        return 1

    try:
        zone_fieldnames, zone_rows, scan_fieldnames, scan_rows = load_metadata_files(
            base_dir, excel_path
        )
    except FileNotFoundError as exc:
        log_message(str(exc), "ERROR")
        return 1

    headers, excel_rows = read_excel_channels(excel_path)
    channel_rows, name_to_new_index, channel_fieldnames = process_excel_rows(
        headers, excel_rows
    )

    zone_order = rebuild_zone_membership_from_excel(excel_rows)
    scan_order = rebuild_scan_membership_from_excel(excel_rows)

    exported_zones = export_zone_csv(
        zone_fieldnames, zone_rows, zone_order, name_to_new_index, channel_rows
    )
    exported_scans = export_scan_csv(
        scan_fieldnames, scan_rows, scan_order, name_to_new_index, channel_rows
    )

    date_prefix = today_str()
    channel_out = base_dir / f"{date_prefix}-Channel.csv"
    zone_out = base_dir / f"{date_prefix}-Zone.csv"
    scan_out = base_dir / f"{date_prefix}-ScanList.csv"

    write_csv_dicts(channel_out, channel_fieldnames, channel_rows)
    write_csv_dicts(zone_out, zone_fieldnames, exported_zones)
    write_csv_dicts(scan_out, scan_fieldnames, exported_scans)

    log_message(f"File generato: {channel_out} ({len(channel_rows)} righe)")
    log_message(f"File generato: {zone_out}")
    log_message(f"File generato: {scan_out}")
    log_message(f"Canali validi re-indicizzati: {len(name_to_new_index)}")
    log_message("Script2 completato con successo")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
