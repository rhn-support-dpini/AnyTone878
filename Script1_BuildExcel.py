#!/usr/bin/env python3
"""Build Master_Codeplug Excel from Channel, Zone and ScanList CSV files."""

from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path
from typing import Dict, List, Set

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils import get_column_letter

from codeplug_utils import (
    COLOR_ANALOG,
    COLOR_DIGITAL,
    CHANNEL_TYPE_ANALOG,
    CHANNEL_TYPE_DIGITAL,
    ensure_unique_channel_names,
    find_input_file,
    is_frequency_column,
    is_single_frequency_column,
    join_pipe_values,
    log_message,
    normalize_frequencies_in_row,
    read_csv_dicts,
    resolve_member_to_name,
    resolve_members_field,
    split_pipe_values,
    today_str,
    write_csv_dicts,
)

WORKSHEET_NAME = "Channels"
MEMBERSHIP_COLUMNS = ("Zone Membership", "ScanList Membership")

ZONE_MEMBER_FIELDS = (
    "Zone Channel Member",
    "A Channel",
    "B Channel",
)

ZONE_MEMBER_FREQ_FIELDS = (
    "Zone Channel Member RX Frequency",
    "Zone Channel Member TX Frequency",
    "A Channel RX Frequency",
    "A Channel TX Frequency",
    "B Channel RX Frequency",
    "B Channel TX Frequency",
)

SCAN_MEMBER_FIELDS = (
    "Scan Channel Member",
    "Priority Channel 1",
    "Priority Channel 2",
)


def convert_zone_rows(
    zone_rows: List[Dict[str, str]],
    old_index_to_name: Dict[int, str],
) -> List[Dict[str, str]]:
    converted: List[Dict[str, str]] = []
    for row in zone_rows:
        zone_name = row.get("Zone Name", "")
        updated = dict(row)
        for field in ZONE_MEMBER_FIELDS:
            if field in updated:
                updated[field] = resolve_members_field(
                    updated[field], old_index_to_name, f"Zone '{zone_name}'/{field}"
                )
        freq_cols = [c for c in updated if is_single_frequency_column(c)]
        updated = normalize_frequencies_in_row(updated, freq_cols, f"Zone '{zone_name}'")
        converted.append(updated)
    return converted


def convert_scan_rows(
    scan_rows: List[Dict[str, str]],
    old_index_to_name: Dict[int, str],
) -> List[Dict[str, str]]:
    converted: List[Dict[str, str]] = []
    for row in scan_rows:
        scan_name = row.get("Scan List Name", "")
        updated = dict(row)
        for field in SCAN_MEMBER_FIELDS:
            if field in updated and updated[field].strip():
                if field == "Scan Channel Member":
                    updated[field] = resolve_members_field(
                        updated[field], old_index_to_name, f"ScanList '{scan_name}'/{field}"
                    )
                else:
                    updated[field] = resolve_member_to_name(
                        updated[field], old_index_to_name, f"ScanList '{scan_name}'/{field}"
                    )
        freq_cols = [c for c in updated if is_single_frequency_column(c)]
        updated = normalize_frequencies_in_row(updated, freq_cols, f"ScanList '{scan_name}'")
        converted.append(updated)
    return converted


def build_membership_maps(
    zone_rows: List[Dict[str, str]],
    scan_rows: List[Dict[str, str]],
    known_channel_names: Set[str],
) -> tuple[Dict[str, List[str]], Dict[str, List[str]]]:
    zone_membership: Dict[str, List[str]] = defaultdict(list)
    scan_membership: Dict[str, List[str]] = defaultdict(list)

    for zone in zone_rows:
        zone_name = zone.get("Zone Name", "").strip()
        if not zone_name:
            continue
        members = split_pipe_values(zone.get("Zone Channel Member", ""))
        for member in members:
            if member not in known_channel_names:
                log_message(
                    f"Riferimento mancante: canale '{member}' in zona '{zone_name}' "
                    "non presente in Channel.csv",
                    "WARNING",
                )
                continue
            if zone_name not in zone_membership[member]:
                zone_membership[member].append(zone_name)

    for scan in scan_rows:
        scan_name = scan.get("Scan List Name", "").strip()
        if not scan_name:
            continue
        members = split_pipe_values(scan.get("Scan Channel Member", ""))
        for member in members:
            if member not in known_channel_names:
                log_message(
                    f"Riferimento mancante: canale '{member}' in scan list '{scan_name}' "
                    "non presente in Channel.csv",
                    "WARNING",
                )
                continue
            if scan_name not in scan_membership[member]:
                scan_membership[member].append(scan_name)

    return zone_membership, scan_membership


def build_excel_rows(
    channels: List[Dict[str, str]],
    channel_fieldnames: List[str],
    zone_membership: Dict[str, List[str]],
    scan_membership: Dict[str, List[str]],
) -> tuple[List[Dict[str, str]], List[str]]:
    excel_columns = list(channel_fieldnames) + list(MEMBERSHIP_COLUMNS)
    rows: List[Dict[str, str]] = []

    freq_cols = [c for c in channel_fieldnames if is_single_frequency_column(c)]

    for row in channels:
        name = row.get("Channel Name", "")
        excel_row = {col: row.get(col, "") for col in channel_fieldnames}
        excel_row = normalize_frequencies_in_row(
            excel_row, freq_cols, f"Channel '{name}'"
        )
        excel_row["Zone Membership"] = join_pipe_values(zone_membership.get(name, []))
        excel_row["ScanList Membership"] = join_pipe_values(scan_membership.get(name, []))
        rows.append(excel_row)

    return rows, excel_columns


def apply_row_formatting(worksheet, row_idx: int, channel_type: str) -> None:
    if channel_type == CHANNEL_TYPE_DIGITAL:
        fill = PatternFill(fill_type="solid", fgColor=COLOR_DIGITAL)
    elif channel_type == CHANNEL_TYPE_ANALOG:
        fill = PatternFill(fill_type="solid", fgColor=COLOR_ANALOG)
    else:
        return

    for col_idx in range(1, worksheet.max_column + 1):
        worksheet.cell(row=row_idx, column=col_idx).fill = fill


def write_excel(path: Path, columns: List[str], rows: List[Dict[str, str]]) -> None:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = WORKSHEET_NAME

    for col_idx, column in enumerate(columns, start=1):
        worksheet.cell(row=1, column=col_idx, value=column)

    freq_col_indexes = {
        idx for idx, col in enumerate(columns, start=1) if is_frequency_column(col)
    }

    for row_idx, row in enumerate(rows, start=2):
        for col_idx, column in enumerate(columns, start=1):
            value = row.get(column, "")
            cell = worksheet.cell(row=row_idx, column=col_idx, value=value)
            if col_idx in freq_col_indexes and value != "":
                try:
                    cell.value = float(str(value))
                except ValueError:
                    cell.value = value
                cell.number_format = "0.00000"

        channel_type = str(row.get("Channel Type", "")).strip()
        apply_row_formatting(worksheet, row_idx, channel_type)

    for col_idx, column in enumerate(columns, start=1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = max(
            12, min(40, len(column) + 2)
        )

    workbook.save(path)


def main() -> int:
    base_dir = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(".")

    try:
        channel_path = find_input_file(base_dir, "Channel.csv")
        zone_path = find_input_file(base_dir, "Zone.csv")
        scan_path = find_input_file(base_dir, "ScanList.csv")
    except FileNotFoundError as exc:
        log_message(str(exc), "ERROR")
        return 1

    log_message("Avvio Script1_BuildExcel.py")

    channel_fieldnames, channels = read_csv_dicts(channel_path)
    if not channels:
        log_message("Channel.csv vuoto o senza dati", "ERROR")
        return 1

    channels, old_index_to_name = ensure_unique_channel_names(channels)
    known_names = set(old_index_to_name.values())

    zone_fieldnames, zone_rows = read_csv_dicts(zone_path)
    scan_fieldnames, scan_rows = read_csv_dicts(scan_path)

    zone_rows = convert_zone_rows(zone_rows, old_index_to_name)
    scan_rows = convert_scan_rows(scan_rows, old_index_to_name)

    zone_membership, scan_membership = build_membership_maps(
        zone_rows, scan_rows, known_names
    )

    excel_rows, excel_columns = build_excel_rows(
        channels, channel_fieldnames, zone_membership, scan_membership
    )

    date_prefix = today_str()
    excel_path = base_dir / f"Master_Codeplug_{date_prefix}.xlsx"
    zone_meta_path = base_dir / f"Master_Codeplug_{date_prefix}_Zone.csv"
    scan_meta_path = base_dir / f"Master_Codeplug_{date_prefix}_ScanList.csv"

    write_excel(excel_path, excel_columns, excel_rows)
    write_csv_dicts(zone_meta_path, zone_fieldnames, zone_rows)
    write_csv_dicts(scan_meta_path, scan_fieldnames, scan_rows)

    log_message(f"Excel generato: {excel_path}")
    log_message(f"Metadati zona salvati: {zone_meta_path}")
    log_message(f"Metadati scan list salvati: {scan_meta_path}")
    log_message(f"Canali elaborati: {len(excel_rows)}")
    log_message("Script1 completato con successo")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
